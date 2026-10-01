from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from conciliacao.cli import main
from conciliacao.exporters import CIELO_HEADERS, QUICKPAY_HEADERS, CieloExporter, QuickPayExporter
from conciliacao.processors import CieloProcessor, QuickPayProcessor
from conciliacao.readers.cielo_reader import ler_cielo
from conciliacao.readers.models import FormatoArquivo
from conciliacao.readers.quickpay_reader import ler_quickpay
from conciliacao.utils.currency import sum_money
from conciliacao.validators.cielo_validator import CieloValidator
from conciliacao.validators.common_validator import valor_original
from conciliacao.validators.quickpay_validator import QuickPayValidator

EXAMPLES = Path("arquivos_exemplo")
CIELO = (
    EXAMPLES / "CIELO VENDA 11.07.2026, 12.07.2026 E 13.07.2026 RECEBIMENTO 13.07.2026 .xls.xlsx"
)
QUICKPAY_HTML = EXAMPLES / "Lista de Transações.20260714094607.xls"
QUICKPAY_XLSX = EXAMPLES / "QUICKPAY VENDA 13.07.2026 RECEBIMENTO 14.07.2026 .xls.xlsx"
CIELO_LEO = (
    EXAMPLES
    / "CIELO VENDA 11.07.2026, 12.07.2026 E 13.07.2026 RECEBIMENTO 13.07.2026 LEO .xls.xlsx"
)
CURRENT_CIELO = EXAMPLES / "CIELO VENDA 14.07.2026 E 15.07.2026 RECEBIMENTO 15.07.2026 .xls.xlsx"


@pytest.mark.local
@pytest.mark.integration
@pytest.mark.skipif(not CIELO.exists(), reason="arquivo Cielo real nao esta disponivel localmente")
def test_reads_local_cielo_example_without_modifying_it():
    result = ler_cielo(CIELO)

    assert result.nome_aba_ou_tabela == "Recebiveis_cielo_detalhe1"
    assert len(result.transacoes) == 52
    assert sum_money(transaction.valor_bruto for transaction in result.transacoes) == Decimal(
        "4356.81"
    )
    assert sum_money(transaction.taxa_original for transaction in result.transacoes) == Decimal(
        "-67.69"
    )
    assert sum_money(transaction.valor_liquido for transaction in result.transacoes) == Decimal(
        "4289.12"
    )

    validation = CieloValidator().validar(result)
    assert validation.valido
    assert validation.totais_calculados["total_bruto"] == Decimal("4356.81")
    assert validation.totais_calculados["total_taxas_originais"] == Decimal("-67.69")


@pytest.mark.local
@pytest.mark.integration
@pytest.mark.skipif(
    not CIELO.exists() or not CIELO_LEO.exists(),
    reason="arquivos Cielo real/LEO nao estao disponiveis localmente",
)
def test_generates_local_cielo_report_structurally_compatible_with_leo(tmp_path):
    leitura = ler_cielo(CIELO)
    validacao = CieloValidator().validar(
        leitura,
        data_inicio="2026-07-10",
        data_fim="2026-07-13",
    )
    relatorio = CieloProcessor().processar(
        leitura,
        validacao,
        data_inicio="2026-07-10",
        data_fim="2026-07-13",
    )
    exportacao = CieloExporter().exportar(relatorio, diretorio_saida=tmp_path)

    generated = load_workbook(exportacao.caminho_saida, data_only=False)
    reference = load_workbook(CIELO_LEO, data_only=False)
    try:
        ws = generated["Planilha1"]
        leo = reference["Planilha1"]
        assert generated.sheetnames == ["Planilha1"]
        assert ws.max_column == 12
        assert [ws.cell(2, column).value for column in range(1, 13)] == CIELO_HEADERS
        assert [leo.cell(2, column).value for column in range(1, 13)] == CIELO_HEADERS
        assert relatorio.resumo.transacoes_incluidas == 52
        assert exportacao.ultima_linha_transacao == 54
        assert exportacao.total_row == 55
        assert exportacao.subtotal_rows == tuple(
            3 + bloco.fim_indice for bloco in relatorio.blocos if bloco.tipo == "CARTAO"
        )
        assert ws["G55"].value == "=SUM(G3:G54)"
        assert ws["H55"].value == "=SUM(H3:H54)"
        assert ws["I55"].value == "=SUM(I3:I54)"
        assert ws["J55"].value == "=1-(I55/G55)"
        assert relatorio.resumo.total_bruto == Decimal("4356.81")
        assert relatorio.resumo.total_taxa == Decimal("-67.69")
        assert relatorio.resumo.total_liquido == Decimal("4289.12")
        assert ws["A1"].fill.fgColor.rgb == leo["A1"].fill.fgColor.rgb
    finally:
        generated.close()
        reference.close()


@pytest.mark.local
@pytest.mark.integration
@pytest.mark.skipif(
    not CURRENT_CIELO.exists(), reason="arquivo Cielo atual 14/15 nao esta disponivel localmente"
)
def test_generates_current_local_cielo_report_when_available(tmp_path):
    leitura = ler_cielo(CURRENT_CIELO)
    validacao = CieloValidator().validar(
        leitura,
        data_inicio="2026-07-14",
        data_fim="2026-07-15",
    )
    relatorio = CieloProcessor().processar(
        leitura,
        validacao,
        data_inicio="2026-07-14",
        data_fim="2026-07-15",
    )
    CieloExporter().exportar(relatorio, diretorio_saida=tmp_path)

    assert relatorio.resumo.transacoes_incluidas == 11
    assert relatorio.resumo.total_bruto == Decimal("532.28")
    assert relatorio.resumo.total_taxa == Decimal("-9.82")
    assert relatorio.resumo.total_liquido == Decimal("522.46")


@pytest.mark.local
@pytest.mark.integration
@pytest.mark.skipif(
    not QUICKPAY_HTML.exists(), reason="arquivo QuickPay HTML real nao esta disponivel localmente"
)
def test_reads_local_quickpay_html_example_without_modifying_it(tmp_path):
    result = ler_quickpay(QUICKPAY_HTML)

    assert result.formato_detectado is FormatoArquivo.HTML
    assert len(result.transacoes) == 2
    assert result.cabecalhos_normalizados[0] == "data da venda"

    validation = QuickPayValidator().validar(result)
    assert validation.valido
    assert validation.totais_calculados["total_recebido_banco"] is None

    output_dir = tmp_path / "quickpay"
    assert (
        main(
            [
                "gerar-quickpay",
                "--arquivo",
                str(QUICKPAY_HTML),
                "--data-inicio",
                "2026-07-13",
                "--data-fim",
                "2026-07-13",
                "--saida",
                str(output_dir),
            ]
        )
        == 0
    )
    assert output_dir.exists()


@pytest.mark.local
@pytest.mark.integration
@pytest.mark.skipif(
    not QUICKPAY_XLSX.exists(), reason="arquivo QuickPay XLSX real nao esta disponivel localmente"
)
def test_reads_local_quickpay_xlsx_and_ignores_legacy_bank_block():
    result = ler_quickpay(QUICKPAY_XLSX)

    assert result.formato_detectado is FormatoArquivo.XLSX
    assert len(result.transacoes) == 2
    assert result.avisos_leitura[0].codigo == "COLUNA_BANCO_QUICKPAY_AUSENTE"

    validation = QuickPayValidator().validar(result)
    assert validation.valido
    assert validation.totais_calculados["total_recebido_banco"] is None


@pytest.mark.local
@pytest.mark.integration
@pytest.mark.skipif(
    not QUICKPAY_HTML.exists(), reason="arquivo QuickPay HTML real nao esta disponivel localmente"
)
def test_generates_local_quickpay_report_from_temporary_prepared_copy(tmp_path):
    raw = ler_quickpay(QUICKPAY_HTML)
    prepared = tmp_path / "quickpay_preparado.xlsx"
    bank_values = [Decimal("121.80"), Decimal("186.30")]
    _write_quickpay_prepared_copy(raw.transacoes, bank_values, prepared)

    leitura = ler_quickpay(prepared)
    validacao = QuickPayValidator().validar(
        leitura,
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
    )
    assert validacao.valido
    relatorio = QuickPayProcessor().processar(
        leitura,
        validacao,
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
    )
    exportacao = QuickPayExporter().exportar(relatorio, diretorio_saida=tmp_path)

    workbook = load_workbook(exportacao.caminho_saida, data_only=False)
    try:
        ws = workbook["Conciliação"]
        assert workbook.sheetnames == ["Conciliação"]
        assert [ws.cell(2, column).value for column in range(1, 18)] == QUICKPAY_HEADERS
        assert relatorio.resumo.transacoes_incluidas == 2
        assert relatorio.resumo.total_bruto == Decimal("320.83")
        assert relatorio.resumo.total_taxa == Decimal("12.72")
        assert relatorio.resumo.total_liquido == Decimal("308.03")
        assert relatorio.resumo.total_bruto_liquido == Decimal("12.80")
        assert relatorio.resumo.total_diferenca_taxa == Decimal("-0.08")
        assert relatorio.resumo.total_recebido_banco == Decimal("308.10")
        assert ws["F5"].value == "=SUM(F3:F4)"
        assert ws["G5"].value == "=SUM(G3:G4)"
        assert ws["H5"].value == "=SUM(H3:H4)"
        assert ws["I5"].value == "=SUM(I3:I4)"
        assert ws["J5"].value == "=SUM(J3:J4)"
        assert ws["L5"].value == "=SUM(L3:L4)"
    finally:
        workbook.close()


def _write_quickpay_prepared_copy(transacoes, bank_values: list[Decimal], path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Lista de Transacoes"
    worksheet.append(["Relatorio artificial derivado de leitura local"])
    worksheet.append(
        [
            "Data da venda",
            "Data de recebimento",
            "Numero de Parcelas",
            "Tipo de pagamento",
            "Valor da Venda",
            "Valor liquido",
            "Taxa",
            "Bandeira",
            "RECEBIDO NO BANCO QUICKPAY",
        ]
    )
    for transacao, bank_value in zip(transacoes, bank_values, strict=True):
        worksheet.append(
            [
                f"{transacao.data_venda:%d/%m/%Y} {transacao.hora_venda:%H:%M}",
                f"{transacao.data_recebimento:%d/%m/%Y}",
                transacao.numero_parcelas,
                valor_original(transacao, "Tipo de pagamento"),
                transacao.valor_bruto,
                transacao.valor_liquido,
                transacao.taxa_original,
                transacao.bandeira,
                bank_value,
            ]
        )
    workbook.save(path)
    workbook.close()
