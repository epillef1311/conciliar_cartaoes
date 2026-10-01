from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conciliacao.exporters.quickpay_exporter import QuickPayExporter
from conciliacao.processors.quickpay_processor import QuickPayProcessor
from conciliacao.readers.quickpay_reader import ler_quickpay
from conciliacao.utils.file_hash import sha256_file
from conciliacao.validators.quickpay_validator import QuickPayValidator

SOURCE = Path.home() / "Downloads" / "QUICKPAY VENDA 30.09.2026 RECEBIMENTO 01.10.2026.xlsx"


@pytest.mark.local
@pytest.mark.integration
@pytest.mark.skipif(not SOURCE.exists(), reason="arquivo QuickPay local não disponível")
def test_september_quickpay_without_bank_receipts(tmp_path):
    original_hash = sha256_file(SOURCE)
    reading = ler_quickpay(SOURCE)
    validation = QuickPayValidator().validar(reading)
    assert validation.valido
    report = QuickPayProcessor().processar(
        reading, validation, data_inicio=date(2026, 9, 30), data_fim=date(2026, 9, 30)
    )
    assert report.resumo.transacoes_incluidas == 16
    assert report.resumo.total_liquido == Decimal("1615.14")
    assert report.resumo.total_recebido_banco is None
    assert report.resumo.diferenca_total_banco_liquido is None
    exported = QuickPayExporter().exportar(report, diretorio_saida=tmp_path)
    workbook = load_workbook(exported.caminho_saida)
    try:
        ws = workbook["Conciliação"]
        assert all(ws.cell(row, 12).value == "NÃO INFORMADO" for row in range(3, 19))
        assert "CONFERÊNCIA BANCÁRIA NÃO REALIZADA" in ws["L19"].value
        assert "Conciliação Bancária" not in workbook.sheetnames
    finally:
        workbook.close()
    assert sha256_file(SOURCE) == original_hash
