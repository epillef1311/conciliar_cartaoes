from __future__ import annotations

import json
from dataclasses import replace
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from typing import cast

from openpyxl import load_workbook
from pydantic import ValidationError

from conciliacao.domain.enums import Modalidade, Operadora
from conciliacao.exporters.quickpay_exporter import SHEET_NAME
from conciliacao.integrations.velo.authentication import StaticTokenProvider
from conciliacao.integrations.velo.categories import CategoriaFiltroVelo
from conciliacao.integrations.velo.testing import FakeVeloTransport, text_response
from conciliacao.processors.quickpay_processor import QuickPayProcessor
from conciliacao.readers.quickpay_reader import ler_quickpay
from conciliacao.validators.models import ResultadoValidacao
from conciliacao.validators.quickpay_validator import QuickPayValidator
from conciliacao.workflow import (
    ReconciliationCommand,
    ReconciliationWorkflow,
    WorkflowMode,
    WorkflowStatus,
    _periodos_consulta_por_categoria,
    _PreparedOperator,
)


def test_workflow_simulado_processa_ambas_operadoras_e_preenche_quickpay(tmp_path):
    api_dir = _api_fixtures(tmp_path)
    output_dir = tmp_path / "output"
    spreadsheets_dir = tmp_path / "planilhas"
    command = ReconciliationCommand(
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
        arquivo_cielo=Path("tests/fixtures/cielo/cielo_valido.xlsx"),
        arquivo_quickpay=Path("tests/fixtures/quickpay/quickpay_valido.xlsx"),
        diretorio_saida=output_dir,
        diretorio_planilhas=spreadsheets_dir,
        modo_simulado=True,
        diretorio_fixtures_api=api_dir,
        salvar_auditoria=True,
        identificador_execucao="20260719_120000_teste1",
    )

    result = ReconciliationWorkflow().executar(command)

    assert result.status_geral is WorkflowStatus.SUCESSO
    assert result.codigo_saida == 0
    assert result.modo is WorkflowMode.SIMULADO
    assert result.resultado_cielo is not None
    assert result.resultado_quickpay is not None
    assert result.resultado_cielo.resumo_matching["resumo"]["conciliados"] == 2
    assert result.resultado_quickpay.resumo_matching["resumo"]["conciliados"] == 2
    assert result.resultado_cielo.hash_preservado
    assert result.resultado_quickpay.hash_preservado
    assert result.resultado_cielo.arquivo_saida is not None
    assert result.resultado_quickpay.arquivo_saida is not None
    assert result.resultado_cielo.arquivo_saida.exists()
    assert result.resultado_quickpay.arquivo_saida.exists()
    expected_spreadsheet_dir = spreadsheets_dir / result.inicio_execucao.date().isoformat()
    assert result.resultado_cielo.arquivo_saida.parent == expected_spreadsheet_dir / "cielo"
    assert result.resultado_quickpay.arquivo_saida.parent == expected_spreadsheet_dir / "quickpay"
    assert (output_dir / "execucoes" / command.identificador_execucao / "resumo.json").exists()
    assert (Path("data/api_raw") / command.identificador_execucao / "metadata.json").exists()

    cielo_workbook = load_workbook(result.resultado_cielo.arquivo_saida, data_only=False)
    try:
        comparison = cielo_workbook["Conciliação Velo"]
        assert cielo_workbook.sheetnames == ["Planilha1", "Conciliação Velo"]
        assert comparison.cell(5, 3).value == 2
        assert comparison.cell(8, 8).value == "CONCILIADO"
    finally:
        cielo_workbook.close()

    workbook = load_workbook(result.resultado_quickpay.arquivo_saida, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        assert worksheet.cell(3, 13).value == 194
        assert worksheet.cell(3, 14).value == 0
        assert worksheet.cell(3, 15).value == "CONCILIADO"
        assert worksheet.cell(4, 13).value == 126.83
        assert worksheet.cell(4, 14).value == 0
        assert worksheet.cell(4, 15).value == "CONCILIADO"
    finally:
        workbook.close()

    combined = _combined_text(result.logs + result.arquivos_de_auditoria)
    assert "Authorization" not in combined
    assert "Bearer" not in combined
    assert "token-teste-nao-real" not in combined


def test_workflow_quickpay_without_bank_receipts_matches_velo(tmp_path):
    source = tmp_path / "quickpay_sem_banco.xlsx"
    workbook = load_workbook("tests/fixtures/quickpay/quickpay_valido.xlsx")
    sheet = workbook.active
    for cell in sheet[2]:
        if cell.value == "RECEBIDO NO BANCO QUICKPAY":
            sheet.delete_cols(cell.column)
            break
    workbook.save(source)
    workbook.close()
    result = ReconciliationWorkflow().executar(
        ReconciliationCommand(
            data_inicio="2026-07-13",
            data_fim="2026-07-13",
            arquivo_quickpay=source,
            diretorio_saida=tmp_path / "output",
            diretorio_planilhas=tmp_path / "planilhas",
            modo_simulado=True,
            diretorio_fixtures_api=_api_fixtures(tmp_path),
            salvar_auditoria=False,
            identificador_execucao="quickpay_sem_banco",
        )
    )
    assert result.status_geral is WorkflowStatus.SUCESSO
    operator = result.resultado_quickpay
    assert operator is not None
    assert operator.hash_preservado
    assert operator.resumo_matching["resumo"]["conciliados"] == 2
    assert operator.resumo_processamento["total_recebido_banco"] is None
    assert operator.resumo_processamento["diferenca_total_banco_liquido"] is None
    assert operator.resumo_processamento["conferencia_bancaria"] == "NAO_REALIZADA"
    generated = load_workbook(operator.arquivo_saida)
    try:
        ws = generated[SHEET_NAME]
        assert ws["L3"].value == "NÃO INFORMADO"
        assert ws["L4"].value == "NÃO INFORMADO"
        assert "CONFERÊNCIA BANCÁRIA NÃO REALIZADA" in ws["L5"].value
        assert ws["O3"].value == "CONCILIADO"
        assert "Conciliação Bancária" not in generated.sheetnames
    finally:
        generated.close()


def test_workflow_somente_cielo_nao_exige_quickpay(tmp_path):
    command = ReconciliationCommand(
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
        arquivo_cielo=Path("tests/fixtures/cielo/cielo_valido.xlsx"),
        diretorio_saida=tmp_path / "output",
        diretorio_planilhas=tmp_path / "planilhas",
        modo_simulado=True,
        diretorio_fixtures_api=_api_fixtures(tmp_path),
        salvar_auditoria=False,
        identificador_execucao="20260719_120000_cielo",
    )

    result = ReconciliationWorkflow().executar(command)

    assert result.status_geral is WorkflowStatus.SUCESSO
    assert result.resultado_cielo is not None
    assert result.resultado_quickpay is None
    assert result.resultado_cielo.arquivo_saida is not None
    assert result.resultado_cielo.arquivo_saida.exists()


def test_workflow_quickpay_invalida_nao_bloqueia_cielo(tmp_path):
    command = ReconciliationCommand(
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
        arquivo_cielo=Path("tests/fixtures/cielo/cielo_valido.xlsx"),
        arquivo_quickpay=Path("tests/fixtures/quickpay/quickpay_sem_tabela.xlsx"),
        diretorio_saida=tmp_path / "output",
        diretorio_planilhas=tmp_path / "planilhas",
        modo_simulado=True,
        diretorio_fixtures_api=_api_fixtures(tmp_path),
        salvar_auditoria=False,
        identificador_execucao="20260719_120000_parcial",
    )

    result = ReconciliationWorkflow().executar(command)

    assert result.status_geral is WorkflowStatus.SUCESSO_PARCIAL
    assert result.codigo_saida == 2
    assert result.resultado_cielo is not None
    assert result.resultado_cielo.status is WorkflowStatus.SUCESSO
    assert result.resultado_cielo.arquivo_saida is not None
    assert result.resultado_quickpay is not None
    assert result.resultado_quickpay.status is WorkflowStatus.FALHA
    assert result.resultado_quickpay.erros


def test_workflow_falha_categoria_api_marca_pendente_sem_bloquear_excel(tmp_path):
    api_dir = _api_fixtures(tmp_path, omit={"cielo_pix.json"})
    command = ReconciliationCommand(
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
        arquivo_cielo=Path("tests/fixtures/cielo/cielo_valido.xlsx"),
        diretorio_saida=tmp_path / "output",
        diretorio_planilhas=tmp_path / "planilhas",
        modo_simulado=True,
        diretorio_fixtures_api=api_dir,
        salvar_auditoria=False,
        identificador_execucao="20260719_120000_apiparcial",
    )

    result = ReconciliationWorkflow().executar(command)

    assert result.status_geral is WorkflowStatus.SUCESSO_PARCIAL
    assert result.resultado_cielo is not None
    assert result.resultado_cielo.status is WorkflowStatus.SUCESSO_PARCIAL
    assert result.resultado_cielo.resumo_matching["resumo"]["conciliados"] == 1
    assert result.resultado_cielo.resumo_matching["resumo"]["pendencias_dados"] == 1
    assert "CIELO_PIX" in result.resultado_api.categorias_com_erro
    assert result.resultado_cielo.arquivo_saida is not None
    assert result.resultado_cielo.arquivo_saida.exists()


def test_workflow_erro_401_nao_gera_relatorio_e_retorna_codigo_autenticacao(tmp_path, monkeypatch):
    monkeypatch.delenv("VELO_BEARER_TOKEN", raising=False)
    transport = FakeVeloTransport(
        {
            "autoCompletarOperadora": text_response(
                '{"erro":"unauthorized"}',
                status_code=401,
            )
        }
    )
    command = ReconciliationCommand(
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
        arquivo_cielo=Path("tests/fixtures/cielo/cielo_valido.xlsx"),
        diretorio_saida=tmp_path / "output",
        diretorio_planilhas=tmp_path / "planilhas",
        identificador_execucao="20260719_120000_auth",
    )

    result = ReconciliationWorkflow(
        transport=transport,
        token_provider=StaticTokenProvider("token-teste-nao-real"),
    ).executar(command)

    assert result.status_geral is WorkflowStatus.FALHA
    assert result.codigo_saida == 4
    assert result.resultado_cielo is None
    assert not list((tmp_path / "planilhas").glob("**/*.xlsx"))
    assert "token-teste-nao-real" not in _combined_text(result.logs)


def test_workflow_command_rejeita_sem_arquivo():
    try:
        ReconciliationCommand(data_inicio="2026-07-13", data_fim="2026-07-13")
    except ValidationError as exc:
        assert "informe pelo menos um arquivo" in str(exc)
    else:
        raise AssertionError("comando sem arquivos deveria falhar")


def test_workflow_queries_debit_category_by_receipt_date():
    leitura = ler_quickpay("tests/fixtures/quickpay/quickpay_valido.xlsx")
    validacao = QuickPayValidator().validar(
        leitura, data_inicio=date(2026, 7, 13), data_fim=date(2026, 7, 13)
    )
    relatorio = QuickPayProcessor().processar(
        leitura,
        validacao,
        data_inicio=date(2026, 7, 13),
        data_fim=date(2026, 7, 13),
    )
    debit_line = replace(
        relatorio.linhas[0],
        transacao=relatorio.linhas[0].transacao.model_copy(
            update={
                "modalidade": Modalidade.DEBITO,
                "data_recebimento": date(2026, 7, 15),
            }
        ),
    )
    prepared = {
        Operadora.QUICKPAY: _PreparedOperator(
            operadora=Operadora.QUICKPAY,
            validacao=cast(ResultadoValidacao, SimpleNamespace(valido=True)),
            relatorio=replace(relatorio, linhas=(debit_line,)),
            categorias=[CategoriaFiltroVelo.QUICKPAY_DEBITO],
        )
    }
    comando = ReconciliationCommand(
        data_inicio="2026-07-13",
        data_fim="2026-07-13",
        arquivo_quickpay=Path("tests/fixtures/quickpay/quickpay_valido.xlsx"),
    )

    periodos = _periodos_consulta_por_categoria(comando, prepared)

    assert periodos[CategoriaFiltroVelo.QUICKPAY_DEBITO] == (
        date(2026, 7, 15),
        date(2026, 7, 15),
    )


def _api_fixtures(tmp_path: Path, *, omit: set[str] | None = None) -> Path:
    omit = omit or set()
    api_dir = tmp_path / "api"
    api_dir.mkdir()
    payloads = {
        "autocomplete_operadoras.json": [
            {"id": 81, "descricao": "Cartao de Credito - Cielo"},
            {"id": 86, "descricao": "Cartao de Debito - Cielo"},
            {"id": 91, "descricao": "Pix - Cielo"},
            {"id": 44, "descricao": "Cartao de Credito - Quickpay"},
            {"id": 51, "descricao": "Cartao de Dedido - Quickpay"},
            {"id": 42, "descricao": "Pix - Quickpay"},
        ],
        "cielo_credito.json": [
            _registro("CC-001", "Cielo", "Cartao de Credito - Cielo", "Visa", "126.83", 81),
        ],
        "cielo_pix.json": [
            _registro("CP-001", "Cielo", "Pix - Cielo", "Pix", "2.90", 91),
        ],
        "quickpay_credito.json": [
            _registro("QC-001", "QuickPay", "Cartao de Credito - Quickpay", "Visa", "126.83", 44),
            _registro(
                "QC-002",
                "QuickPay",
                "Cartao de Credito - Quickpay",
                "Mastercard",
                "194.00",
                44,
            ),
        ],
    }
    for filename, payload in payloads.items():
        if filename in omit:
            continue
        (api_dir / filename).write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    return api_dir


def _registro(
    ident: str,
    operadora: str,
    forma: str,
    tipo: str,
    valor: str,
    forma_id: int,
) -> dict[str, object]:
    return {
        "contaPacotePagamentoUnicoId": ident,
        "operadoraId": forma_id,
        "operadora": operadora,
        "cadastroCaixaId": f"CX-{ident}",
        "formaRecebimentoId": forma_id,
        "formaRecebimento": forma,
        "tipoCartao": tipo,
        "valor": valor,
        "valorTaxaCartao": "0.00",
        "dataVencimento": "2026-07-14",
        "dataCadastro": "2026-07-13",
        "campoExtra": "anonimo",
    }


def _combined_text(paths: list[Path]) -> str:
    chunks = []
    for path in paths:
        if path.exists() and path.is_file():
            chunks.append(path.read_text(encoding="utf-8"))
    return "\n".join(chunks)
