from __future__ import annotations

from dataclasses import replace
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from conciliacao.domain.enums import Modalidade
from conciliacao.exporters import QUICKPAY_HEADERS, QuickPayExporter
from conciliacao.exporters.quickpay_exporter import SHEET_NAME, STATUS_PENDENTE
from conciliacao.processors.quickpay_processor import QuickPayProcessor
from conciliacao.readers.quickpay_reader import ler_quickpay
from conciliacao.validators.quickpay_validator import QuickPayValidator


def _relatorio_fixture():
    leitura = ler_quickpay("tests/fixtures/quickpay/quickpay_valido.xlsx")
    validacao = QuickPayValidator().validar(
        leitura,
        data_inicio=date(2026, 7, 13),
        data_fim=date(2026, 7, 13),
    )
    assert validacao.valido
    return QuickPayProcessor().processar(
        leitura,
        validacao,
        data_inicio=date(2026, 7, 13),
        data_fim=date(2026, 7, 13),
    )


def test_exports_quickpay_sheet_headers_formulas_total_and_status(tmp_path):
    relatorio = _relatorio_fixture()
    result = QuickPayExporter().exportar(relatorio, diretorio_saida=tmp_path)

    assert result.validacao_saida_ok
    workbook = load_workbook(result.caminho_saida, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        assert workbook.sheetnames == [SHEET_NAME]
        assert worksheet.max_column == 17
        assert [worksheet.cell(2, column).value for column in range(1, 18)] == QUICKPAY_HEADERS
        assert worksheet.cell(3, 9).value == "=F3-H3"
        assert worksheet.cell(3, 10).value == "=G3-I3"
        assert worksheet.cell(3, 11).value == '=IF(F3=0,"",1-(H3/F3))'
        assert worksheet.cell(5, 6).value == "=SUM(F3:F4)"
        assert worksheet.cell(5, 7).value == "=SUM(G3:G4)"
        assert worksheet.cell(5, 8).value == "=SUM(H3:H4)"
        assert worksheet.cell(5, 9).value == "=SUM(I3:I4)"
        assert worksheet.cell(5, 10).value == "=SUM(J3:J4)"
        assert worksheet.cell(5, 12).value == "=SUM(L3:L4)"
        assert worksheet.cell(5, 13).value == "=SUM(M3:M4)"
        assert worksheet.cell(5, 14).value == "=SUM(N3:N4)"
        assert worksheet.cell(3, 16).value is None
        assert worksheet.cell(3, 17).value is None
        assert worksheet.cell(4, 16).value == "=SUM(F3:F4)"
        assert worksheet.cell(4, 17).value == "=SUM(H3:H4)"
        assert worksheet.cell(3, 13).value is None
        assert worksheet.cell(3, 14).value is None
        assert worksheet.cell(3, 15).value == STATUS_PENDENTE
        assert result.total_row == 5
        assert result.ultima_linha_transacao == 4
    finally:
        workbook.close()


def test_export_organizes_by_brand_and_uses_cielo_colors(tmp_path):
    relatorio = _relatorio_fixture()
    debit_line = replace(
        relatorio.linhas[0],
        transacao=relatorio.linhas[0].transacao.model_copy(
            update={"modalidade": Modalidade.DEBITO}
        ),
    )
    relatorio_com_debito = replace(relatorio, linhas=(debit_line, relatorio.linhas[1]))

    result = QuickPayExporter().exportar(relatorio_com_debito, diretorio_saida=tmp_path)

    workbook = load_workbook(result.caminho_saida, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        assert worksheet.cell(3, 1).fill.fgColor.rgb == "FF33CCCC"
        assert worksheet.cell(4, 1).fill.fgColor.rgb == "FFC0C0C0"
    finally:
        workbook.close()


def test_partial_bank_receipts_preserve_known_values_without_false_total(tmp_path):
    reading = ler_quickpay("tests/fixtures/quickpay/quickpay_valido.xlsx")
    tx = reading.transacoes[0]
    raw = dict(tx.dados_originais)
    raw.pop("recebido_no_banco_quickpay_normalizado")
    raw["valores"] = {
        key: value
        for key, value in raw["valores"].items()
        if not key.startswith("RECEBIDO NO BANCO QUICKPAY")
    }
    reading.transacoes[0] = tx.model_copy(update={"dados_originais": raw})
    validation = QuickPayValidator().validar(reading)
    assert validation.valido
    report = QuickPayProcessor().processar(
        reading, validation, data_inicio=date(2026, 7, 13), data_fim=date(2026, 7, 13)
    )
    assert report.resumo.total_recebido_banco is None
    assert report.resumo.diferenca_total_banco_liquido is None
    result = QuickPayExporter().exportar(report, diretorio_saida=tmp_path)
    wb = load_workbook(result.caminho_saida)
    try:
        values = [wb[SHEET_NAME].cell(row, 12).value for row in (3, 4)]
        assert "NÃO INFORMADO" in values
        assert any(isinstance(value, (int, float)) for value in values)
        assert "CONFERÊNCIA BANCÁRIA NÃO REALIZADA" in wb[SHEET_NAME]["L5"].value
    finally:
        wb.close()


def test_export_preserves_quickpay_bank_values_and_one_cent_differences(tmp_path):
    relatorio = _relatorio_fixture()
    result = QuickPayExporter().exportar(relatorio, diretorio_saida=tmp_path)

    workbook = load_workbook(result.caminho_saida, data_only=False)
    try:
        worksheet = workbook[SHEET_NAME]
        assert Decimal(str(worksheet.cell(3, 12).value)) == relatorio.linhas[0].recebido_banco
        assert Decimal(str(worksheet.cell(4, 12).value)) == relatorio.linhas[1].recebido_banco
        assert relatorio.resumo.total_bruto == Decimal("320.83")
        assert relatorio.resumo.total_taxa == Decimal("12.72")
        assert relatorio.resumo.total_liquido == Decimal("308.03")
        assert relatorio.resumo.total_bruto_liquido == Decimal("12.80")
        assert relatorio.resumo.total_diferenca_taxa == Decimal("-0.08")
        assert relatorio.resumo.total_recebido_banco == Decimal("308.10")
        assert worksheet.cell(1, 1).value.startswith("VENDAS QUICKPAY FRIGORIFICO CANDEIAS")
    finally:
        workbook.close()


def test_quickpay_export_does_not_overwrite_silently_and_can_overwrite_explicitly(tmp_path):
    relatorio = _relatorio_fixture()
    first = QuickPayExporter().exportar(relatorio, diretorio_saida=tmp_path)
    second = QuickPayExporter().exportar(relatorio, diretorio_saida=tmp_path)
    explicit = QuickPayExporter().exportar(
        relatorio,
        diretorio_saida=tmp_path,
        sobrescrever=True,
    )

    assert first.caminho_saida.exists()
    assert second.caminho_saida.exists()
    assert second.caminho_saida != first.caminho_saida
    assert explicit.caminho_saida == first.caminho_saida
