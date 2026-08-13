from __future__ import annotations

from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from conciliacao.exporters import CIELO_HEADERS, CieloExporter
from conciliacao.processors.cielo_processor import CieloProcessor
from conciliacao.readers.cielo_reader import ler_cielo
from conciliacao.validators.cielo_validator import CieloValidator


def _relatorio_fixture():
    leitura = ler_cielo("tests/fixtures/cielo/cielo_valido.xlsx")
    validacao = CieloValidator().validar(
        leitura,
        data_inicio=date(2026, 7, 13),
        data_fim=date(2026, 7, 13),
    )
    assert validacao.valido
    return CieloProcessor().processar(
        leitura,
        validacao,
        data_inicio=date(2026, 7, 13),
        data_fim=date(2026, 7, 13),
    )


def test_exports_cielo_sheet_headers_formulas_subtotals_and_total(tmp_path):
    relatorio = _relatorio_fixture()
    result = CieloExporter().exportar(relatorio, diretorio_saida=tmp_path)

    assert result.validacao_saida_ok
    workbook = load_workbook(result.caminho_saida, data_only=False)
    try:
        worksheet = workbook["Planilha1"]
        assert workbook.sheetnames == ["Planilha1"]
        assert worksheet.max_column == 12
        assert [worksheet.cell(2, column).value for column in range(1, 13)] == CIELO_HEADERS
        assert worksheet.cell(3, 10).value == "=1-(I3/G3)"
        assert worksheet.cell(3, 11).value == "=SUM(G3:G3)"
        assert worksheet.cell(4, 11).value is None
        assert worksheet.cell(5, 7).value == "=SUM(G3:G4)"
        assert worksheet.cell(5, 8).value == "=SUM(H3:H4)"
        assert worksheet.cell(5, 9).value == "=SUM(I3:I4)"
        assert worksheet.cell(5, 10).value == "=1-(I5/G5)"
        assert result.subtotal_rows == (3,)
        assert worksheet.cell(3, 1).fill.fgColor.rgb == "FFC0C0C0"
        assert worksheet.cell(4, 1).fill.fgColor.rgb == "FFFFFF99"
        assert result.total_row == 5
    finally:
        workbook.close()


def test_export_preserves_negative_fee_and_reopens_values(tmp_path):
    relatorio = _relatorio_fixture()
    result = CieloExporter().exportar(relatorio, diretorio_saida=tmp_path)

    workbook = load_workbook(result.caminho_saida, data_only=False)
    try:
        worksheet = workbook["Planilha1"]
        assert Decimal(str(worksheet.cell(3, 8).value)) == Decimal("-5.03")
        assert Decimal(str(worksheet.cell(4, 8).value)) == Decimal("-0.02")
        assert worksheet.cell(1, 1).value.startswith("VENDAS CIELO FRIGORIFICO CANDEIAS")
    finally:
        workbook.close()


def test_export_does_not_overwrite_silently_and_can_overwrite_explicitly(tmp_path):
    relatorio = _relatorio_fixture()
    first = CieloExporter().exportar(relatorio, diretorio_saida=tmp_path)
    second = CieloExporter().exportar(relatorio, diretorio_saida=tmp_path)
    explicit = CieloExporter().exportar(relatorio, diretorio_saida=tmp_path, sobrescrever=True)

    assert first.caminho_saida.exists()
    assert second.caminho_saida.exists()
    assert second.caminho_saida != first.caminho_saida
    assert explicit.caminho_saida == first.caminho_saida
