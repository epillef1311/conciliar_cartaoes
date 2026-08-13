"""Leitores somente leitura para exportacoes QuickPay em XLSX ou HTML."""

from collections.abc import Sequence
from datetime import date, time
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from conciliacao.domain.enums import Operadora, OrigemRegistro
from conciliacao.domain.exceptions import DataError, MonetarioError
from conciliacao.domain.models import TransacaoOperadora
from conciliacao.readers.common import (
    cabecalhos_normalizados,
    cabecalhos_originais,
    converter_data,
    converter_moeda,
    dados_originais,
    data_e_hora_venda,
    leitura_somente_leitura,
    linha_total_ou_vazia,
    modalidade_por_texto,
    valor_obrigatorio,
    valor_opcional,
)
from conciliacao.readers.exceptions import (
    DataLeituraInvalidaError,
    FormatoNaoSuportadoError,
    ValorMonetarioInvalidoError,
)
from conciliacao.readers.format_detector import detectar_formato
from conciliacao.readers.headers import (
    CabecalhoEncontrado,
    LinhaParaBusca,
    RegraCabecalho,
    localizar_cabecalho,
)
from conciliacao.readers.models import FormatoArquivo, ResultadoLeitura, aviso_coluna_banco_ausente
from conciliacao.utils.currency import parse_money
from conciliacao.utils.text import normalize_text

QUICKPAY_REGRAS = (
    RegraCabecalho("data_venda", ("Data da venda",)),
    RegraCabecalho("data_recebimento", ("Data de recebimento",)),
    RegraCabecalho("numero_parcelas", ("Número de Parcelas",)),
    RegraCabecalho("tipo_pagamento", ("Tipo de pagamento",)),
    RegraCabecalho("valor_venda", ("Valor da Venda",)),
    RegraCabecalho("valor_liquido", ("Valor líquido",)),
    RegraCabecalho("taxa", ("Taxa",)),
    RegraCabecalho("bandeira", ("Bandeira",)),
)
_BANCO_HEADER = "recebido no banco quickpay"


def ler_quickpay(path: str | Path) -> ResultadoLeitura:
    """Le QuickPay XLSX ou HTML e nunca valida a coluna bancaria nesta etapa."""
    file_path = Path(path)
    with leitura_somente_leitura(file_path) as snapshot:
        formato = detectar_formato(file_path)
        if formato is FormatoArquivo.XLSX:
            table = _ler_tabela_xlsx(file_path)
        elif formato is FormatoArquivo.HTML:
            table = _ler_tabela_html(file_path)
        else:
            raise FormatoNaoSuportadoError(
                f"o leitor QuickPay nao suporta leitura de {formato.value} nesta etapa"
            )

        transacoes = _extrair_transacoes(
            values_rows=table.rows,
            first_data_row=table.header.numero_linha + 1,
            headers=table.headers,
            indices=table.header.indices,
            arquivo=file_path,
            aba=table.header.local,
        )
        banco_index = _indice_banco(table.headers)
        avisos = []
        if banco_index is None:
            avisos.append(aviso_coluna_banco_ausente(arquivo=file_path, aba=table.header.local))

    return ResultadoLeitura(
        caminho_arquivo=file_path,
        formato_detectado=formato,
        nome_aba_ou_tabela=table.header.local,
        cabecalhos_originais=table.headers,
        cabecalhos_normalizados=cabecalhos_normalizados(table.header.valores),
        transacoes=transacoes,
        avisos_leitura=avisos,
        metadados_arquivo=snapshot.metadados,
        hash_sha256=snapshot.hash_sha256,
    )


class _TabelaLida:
    def __init__(
        self,
        *,
        header: CabecalhoEncontrado,
        headers: list[str],
        rows: Sequence[Sequence[object]],
    ) -> None:
        self.header = header
        self.headers = headers
        self.rows = rows


def _ler_tabela_xlsx(path: Path) -> _TabelaLida:
    workbook = load_workbook(
        path, read_only=True, data_only=False, keep_links=False, keep_vba=False
    )
    try:
        lines: list[LinhaParaBusca] = []
        for worksheet in workbook.worksheets:
            for number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if number > 200:
                    break
                lines.append(LinhaParaBusca(worksheet.title, number, row))
        header = localizar_cabecalho(lines, QUICKPAY_REGRAS, arquivo=str(path))
        worksheet = workbook[header.local]
        headers = cabecalhos_originais(header.valores)
        rows = [
            row for row in worksheet.iter_rows(min_row=header.numero_linha + 1, values_only=True)
        ]
        return _TabelaLida(header=header, headers=headers, rows=rows)
    finally:
        workbook.close()


def _ler_tabela_html(path: Path) -> _TabelaLida:
    parser = _ParserTabelasHtml()
    try:
        parser.feed(path.read_text(encoding="utf-8"))
        parser.close()
    except UnicodeDecodeError as exc:
        raise FormatoNaoSuportadoError(f"HTML QuickPay nao esta em UTF-8: {path}") from exc

    lines: list[LinhaParaBusca] = []
    for index, table in enumerate(parser.tables, start=1):
        for number, row in enumerate(table, start=1):
            lines.append(LinhaParaBusca(f"tabela_html_{index}", number, row))
    header = localizar_cabecalho(lines, QUICKPAY_REGRAS, arquivo=str(path))
    table_index = int(header.local.rsplit("_", maxsplit=1)[1]) - 1
    headers = cabecalhos_originais(header.valores)
    return _TabelaLida(
        header=header, headers=headers, rows=parser.tables[table_index][header.numero_linha :]
    )


def _extrair_transacoes(
    *,
    values_rows: Sequence[Sequence[object]],
    first_data_row: int,
    headers: Sequence[str],
    indices: dict[str, int],
    arquivo: Path,
    aba: str,
) -> list[TransacaoOperadora]:
    transacoes: list[TransacaoOperadora] = []
    banco_index = _indice_banco(headers)
    for row_number, values in enumerate(values_rows, start=first_data_row):
        if linha_total_ou_vazia(values) or _linha_total_quickpay(values, indices):
            break
        transacoes.append(
            _criar_transacao(
                values=values,
                row_number=row_number,
                headers=headers,
                indices=indices,
                banco_index=banco_index,
                arquivo=arquivo,
                aba=aba,
            )
        )
    return transacoes


def _linha_total_quickpay(values: Sequence[object], indices: dict[str, int]) -> bool:
    """Distingue o total formulaico do exemplo de uma transacao incompleta."""
    sale_value = valor_opcional(values, indices["data_venda"])
    formulas = [value for value in values if isinstance(value, str) and value.startswith("=")]
    return sale_value is None and bool(formulas)


def _criar_transacao(
    *,
    values: Sequence[object],
    row_number: int,
    headers: Sequence[str],
    indices: dict[str, int],
    banco_index: int | None,
    arquivo: Path,
    aba: str,
) -> TransacaoOperadora:
    sale_value = valor_obrigatorio(
        values, indices, "data_venda", arquivo=arquivo, aba=aba, linha=row_number
    )
    sale_date, sale_time = _data_e_hora_venda(
        sale_value,
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["data_venda"],
    )
    receipt = converter_data(
        valor_obrigatorio(
            values, indices, "data_recebimento", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["data_recebimento"],
        cabecalho="Data de recebimento",
    )
    parcelas_value = valor_opcional(values, indices["numero_parcelas"])
    parcelas = (
        None
        if parcelas_value is None
        else _parcelas(
            parcelas_value,
            arquivo=arquivo,
            aba=aba,
            linha=row_number,
        )
    )
    bruto = converter_moeda(
        valor_obrigatorio(
            values, indices, "valor_venda", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["valor_venda"],
        cabecalho="Valor da Venda",
    )
    taxa = converter_moeda(
        valor_obrigatorio(values, indices, "taxa", arquivo=arquivo, aba=aba, linha=row_number),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["taxa"],
        cabecalho="Taxa",
    )
    liquido = converter_moeda(
        valor_obrigatorio(
            values, indices, "valor_liquido", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["valor_liquido"],
        cabecalho="Valor líquido",
    )
    tipo = valor_obrigatorio(
        values, indices, "tipo_pagamento", arquivo=arquivo, aba=aba, linha=row_number
    )
    raw_data, cell_map = dados_originais(headers=headers, values=values, row_number=row_number)
    banco_value = valor_opcional(values, banco_index)
    if banco_value is not None and banco_index is not None:
        try:
            raw_data["recebido_no_banco_quickpay_normalizado"] = converter_moeda(
                banco_value,
                arquivo=arquivo,
                aba=aba,
                linha=row_number,
                coluna=banco_index,
                cabecalho="RECEBIDO NO BANCO QUICKPAY",
            )
        except ValorMonetarioInvalidoError:
            raw_data["recebido_no_banco_quickpay_invalido"] = banco_value
    return TransacaoOperadora(
        operadora=Operadora.QUICKPAY,
        modalidade=modalidade_por_texto(tipo),
        bandeira=str(
            valor_obrigatorio(
                values, indices, "bandeira", arquivo=arquivo, aba=aba, linha=row_number
            )
        ),
        data_venda=sale_date,
        hora_venda=sale_time,
        data_recebimento=receipt,
        numero_parcelas=parcelas,
        valor_bruto=bruto,
        valor_liquido=liquido,
        taxa_normalizada=abs(taxa),
        taxa_original=taxa,
        origem_arquivo=OrigemRegistro.QUICKPAY,
        linha_original=row_number,
        celulas_origem=cell_map,
        dados_originais=raw_data,
    )


def _indice_banco(headers: Sequence[str]) -> int | None:
    for index, header in enumerate(headers):
        if normalize_text(header).comparavel == _BANCO_HEADER:
            return index
    return None


def _data_e_hora_venda(
    value: object, *, arquivo: Path, aba: str, linha: int, coluna: int
) -> tuple[date, time | None]:
    try:
        return data_e_hora_venda(value, None)
    except DataError as exc:
        raise DataLeituraInvalidaError(
            "data ou hora da venda invalida",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            coluna=get_column_letter(coluna + 1),
            celula=f"{get_column_letter(coluna + 1)}{linha}",
            cabecalho_esperado="Data da venda",
            valor_recebido=value,
        ) from exc


def _parcelas(value: object, *, arquivo: Path, aba: str, linha: int) -> int:
    try:
        parcelas = int(parse_money(value))
    except MonetarioError as exc:
        raise DataLeituraInvalidaError(
            "numero de parcelas invalido",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            valor_recebido=value,
        ) from exc
    if parcelas < 1:
        raise DataLeituraInvalidaError(
            "numero de parcelas deve ser maior que zero",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            valor_recebido=value,
        )
    return parcelas


class _ParserTabelasHtml(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag == "table" and self._table is None:
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in {"th", "td"} and self._row is not None:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"th", "td"} and self._cell is not None and self._row is not None:
            self._row.append("".join(self._cell).strip(" \t\r\n"))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            self.tables.append(self._table)
            self._table = None
