"""Leitor somente leitura do relatorio bruto da Cielo em XLSX."""

from collections.abc import Sequence
from datetime import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from conciliacao.domain.enums import Operadora, OrigemRegistro
from conciliacao.domain.exceptions import DataError, MonetarioError
from conciliacao.domain.models import TransacaoOperadora
from conciliacao.readers.common import (
    cabecalhos_normalizados,
    cabecalhos_originais,
    converter_data,
    converter_hora,
    converter_moeda,
    dados_originais,
    leitura_somente_leitura,
    linha_total_ou_vazia,
    modalidade_por_texto,
    valor_obrigatorio,
    valor_opcional,
)
from conciliacao.readers.exceptions import (
    DataLeituraInvalidaError,
    FormatoNaoSuportadoError,
    LinhaIncompletaError,
)
from conciliacao.readers.format_detector import detectar_formato
from conciliacao.readers.headers import LinhaParaBusca, RegraCabecalho, localizar_cabecalho
from conciliacao.readers.models import FormatoArquivo, ResultadoLeitura
from conciliacao.utils.currency import parse_money
from conciliacao.utils.text import normalize_text

CIELO_REGRAS = (
    RegraCabecalho("data_pagamento", ("Data de pagamento", "Data do pagamento")),
    RegraCabecalho("tipo_lancamento", ("Tipo de lançamento",)),
    RegraCabecalho("forma_pagamento", ("Forma de pagamento",)),
    RegraCabecalho("valor_bruto", ("Valor bruto",)),
    RegraCabecalho("taxa_tarifa", ("Taxa/tarifa", "Taxa", "Tarifa")),
    RegraCabecalho("valor_liquido", ("Valor líquido",)),
    RegraCabecalho("data_venda", ("Data da venda",)),
)


def ler_cielo(path: str | Path) -> ResultadoLeitura:
    """Le o relatorio Cielo sem salvar workbook ou criar arquivos auxiliares."""
    file_path = Path(path)
    with leitura_somente_leitura(file_path) as snapshot:
        formato = detectar_formato(file_path)
        if formato is not FormatoArquivo.XLSX:
            raise FormatoNaoSuportadoError(
                f"o leitor Cielo requer XLSX; formato detectado: {formato.value}"
            )

        # O nome do arquivo pode terminar em .xls mesmo quando a assinatura indica
        # um XLSX. Abrir pelo fluxo binario evita que o openpyxl use a extensao para
        # recusar um conteudo que ja foi validado por detectar_formato().
        with file_path.open("rb") as source:
            workbook = load_workbook(
                source,
                read_only=True,
                data_only=False,
                keep_links=False,
                keep_vba=False,
            )
            try:
                _recalcular_dimensoes(workbook)
                linhas = _linhas_para_busca(workbook)
                header = localizar_cabecalho(linhas, CIELO_REGRAS, arquivo=str(file_path))
                worksheet = workbook[header.local]
                headers = cabecalhos_originais(header.valores)
                totalizador = _extrair_totalizador(worksheet, header.numero_linha)
                transactions = _extrair_transacoes(
                    worksheet=worksheet,
                    header_row=header.numero_linha,
                    headers=headers,
                    indices=header.indices,
                    totalizador=totalizador,
                    arquivo=file_path,
                )
            finally:
                workbook.close()

    return ResultadoLeitura(
        caminho_arquivo=file_path,
        formato_detectado=formato,
        nome_aba_ou_tabela=header.local,
        cabecalhos_originais=headers,
        cabecalhos_normalizados=cabecalhos_normalizados(header.valores),
        transacoes=transactions,
        metadados_arquivo=snapshot.metadados,
        hash_sha256=snapshot.hash_sha256,
    )


def _linhas_para_busca(workbook: Workbook) -> list[LinhaParaBusca]:
    linhas: list[LinhaParaBusca] = []
    for worksheet in workbook.worksheets:
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_number > 200:
                break
            linhas.append(LinhaParaBusca(worksheet.title, row_number, row))
    return linhas


def _recalcular_dimensoes(workbook: Workbook) -> None:
    """Ignora dimensoes declaradas incorretamente em relatorios exportados pela Cielo."""
    for worksheet in workbook.worksheets:
        worksheet.reset_dimensions()


def _extrair_transacoes(
    *,
    worksheet: Worksheet,
    header_row: int,
    headers: Sequence[str],
    indices: dict[str, int],
    totalizador: dict[str, object] | None,
    arquivo: Path,
) -> list[TransacaoOperadora]:
    transactions: list[TransacaoOperadora] = []
    optional_indices = _indices_opcionais(headers)
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if linha_total_ou_vazia(row):
            break
        transactions.append(
            _criar_transacao(
                values=row,
                row_number=row_number,
                headers=headers,
                indices=indices,
                optional_indices=optional_indices,
                totalizador=totalizador,
                arquivo=arquivo,
                aba=worksheet.title,
            )
        )
    return transactions


def _criar_transacao(
    *,
    values: Sequence[object],
    row_number: int,
    headers: Sequence[str],
    indices: dict[str, int],
    optional_indices: dict[str, int],
    totalizador: dict[str, object] | None,
    arquivo: Path,
    aba: str,
) -> TransacaoOperadora:
    pagamento = converter_data(
        valor_obrigatorio(
            values, indices, "data_pagamento", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["data_pagamento"],
        cabecalho="Data de pagamento",
    )
    venda = converter_data(
        valor_obrigatorio(
            values, indices, "data_venda", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["data_venda"],
        cabecalho="Data da venda",
    )
    hora = _converter_hora_cielo(
        valor_opcional(values, optional_indices.get("hora_venda")),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=optional_indices.get("hora_venda"),
    )
    tipo = valor_obrigatorio(
        values, indices, "tipo_lancamento", arquivo=arquivo, aba=aba, linha=row_number
    )
    bruto = converter_moeda(
        valor_obrigatorio(
            values, indices, "valor_bruto", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["valor_bruto"],
        cabecalho="Valor bruto",
    )
    taxa_original = converter_moeda(
        valor_obrigatorio(
            values, indices, "taxa_tarifa", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["taxa_tarifa"],
        cabecalho="Taxa/tarifa",
    )
    liquido = converter_moeda(
        valor_obrigatorio(
            values, indices, "valor_liquido", arquivo=arquivo, aba=aba, linha=row_number
        ),
        arquivo=arquivo,
        aba=aba,
        linha=row_number,
        coluna=indices["valor_liquido"],
        cabecalho="Valor líquido",
    )
    raw_data, cell_map = dados_originais(headers=headers, values=values, row_number=row_number)
    raw_data["campos_cielo"] = {
        name: valor_opcional(values, index) for name, index in optional_indices.items()
    }
    raw_data["data_pagamento"] = pagamento
    if totalizador is not None:
        raw_data["totalizador_cielo"] = totalizador
    identificador = _primeiro_identificador(values, optional_indices)
    parcelas = _parcelas_opcionais(values, optional_indices, arquivo, aba, row_number)
    return TransacaoOperadora(
        identificador_origem=identificador,
        operadora=Operadora.CIELO,
        modalidade=modalidade_por_texto(tipo),
        bandeira=_texto(valor_opcional(values, optional_indices.get("bandeira"))),
        data_venda=venda,
        hora_venda=hora,
        data_recebimento=pagamento,
        numero_parcelas=parcelas,
        valor_bruto=bruto,
        valor_liquido=liquido,
        taxa_normalizada=abs(taxa_original),
        taxa_original=taxa_original,
        origem_arquivo=OrigemRegistro.CIELO,
        linha_original=row_number,
        celulas_origem=cell_map,
        dados_originais=raw_data,
    )


def _indices_opcionais(headers: Sequence[str]) -> dict[str, int]:
    aliases = {
        "hora_venda": {"hora da venda"},
        "bandeira": {"bandeira"},
        "codigo_autorizacao": {"codigo da autorizacao"},
        "nsu_doc": {"nsu/doc", "nsu", "doc"},
        "codigo_venda": {"codigo da venda", "codigo de venda"},
        "tid": {"tid"},
        "id_pix": {"id pix", "txid", "id de pagamento pix"},
        "numero_parcela": {"numero da parcela"},
    }
    found: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = normalize_text(header).comparavel
        for field, candidates in aliases.items():
            if field not in found and normalized in candidates:
                found[field] = index
    return found


def _extrair_totalizador(worksheet: Worksheet, header_row: int) -> dict[str, object] | None:
    rows = list(worksheet.iter_rows(min_row=1, max_row=max(header_row - 1, 1), values_only=True))
    for index, row in enumerate(rows[:-1]):
        normalized = [
            normalize_text(value).comparavel if isinstance(value, str) else "" for value in row
        ]
        if not {
            "quantidade de lancamentos",
            "valor bruto",
            "taxa/tarifa",
            "valor liquido",
        }.issubset(set(normalized)):
            continue
        values = rows[index + 1]
        positions = {name: normalized.index(name) for name in set(normalized) if name}
        return {
            "quantidade_lancamentos": _valor_totalizador(
                values, positions, "quantidade de lancamentos"
            ),
            "valor_bruto": _valor_totalizador(values, positions, "valor bruto"),
            "taxa_tarifa": _valor_totalizador(values, positions, "taxa/tarifa"),
            "valor_liquido": _valor_totalizador(values, positions, "valor liquido"),
        }
    return None


def _valor_totalizador(
    values: Sequence[object], positions: dict[str, int], header: str
) -> object | None:
    index = positions.get(header)
    if index is None or index >= len(values):
        return None
    return values[index]


def _converter_hora_cielo(
    value: object | None, *, arquivo: Path, aba: str, linha: int, coluna: int | None
) -> time | None:
    try:
        return converter_hora(value)
    except DataError as exc:
        if coluna is None:
            raise
        letter = get_column_letter(coluna + 1)
        raise DataLeituraInvalidaError(
            "hora da venda invalida",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            coluna=letter,
            celula=f"{letter}{linha}",
            cabecalho_esperado="Hora da venda",
            valor_recebido=value,
        ) from exc


def _primeiro_identificador(values: Sequence[object], indices: dict[str, int]) -> str | None:
    for field in ("codigo_venda", "codigo_autorizacao", "nsu_doc", "tid", "id_pix"):
        value = valor_opcional(values, indices.get(field))
        if value is not None:
            return str(value)
    return None


def _parcelas_opcionais(
    values: Sequence[object], indices: dict[str, int], arquivo: Path, aba: str, linha: int
) -> int | None:
    value = valor_opcional(values, indices.get("numero_parcela"))
    if value is None:
        return None
    try:
        return int(parse_money(value))
    except MonetarioError as exc:
        raise LinhaIncompletaError(
            "numero de parcela invalido",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            valor_recebido=value,
        ) from exc


def _texto(value: object | None) -> str | None:
    return str(value) if value is not None else None
