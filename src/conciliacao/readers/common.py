"""Funcoes compartilhadas pelos leitores sem qualquer escrita em disco."""

from collections.abc import Iterator, Sequence
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl.utils import get_column_letter

from conciliacao.domain.enums import Modalidade
from conciliacao.domain.exceptions import DataError, MonetarioError
from conciliacao.readers.exceptions import (
    ArquivoCorrompidoError,
    DataLeituraInvalidaError,
    LinhaIncompletaError,
    ValorMonetarioInvalidoError,
)
from conciliacao.readers.models import MetadadosArquivo
from conciliacao.utils.currency import parse_money
from conciliacao.utils.dates import parse_date
from conciliacao.utils.file_hash import sha256_file
from conciliacao.utils.text import normalize_text


@dataclass(frozen=True, slots=True)
class InstantaneoArquivo:
    tamanho_bytes: int
    modificado_em_ns: int
    hash_sha256: str

    @property
    def metadados(self) -> MetadadosArquivo:
        return MetadadosArquivo(
            tamanho_bytes=self.tamanho_bytes,
            modificado_em_ns=self.modificado_em_ns,
        )


@contextmanager
def leitura_somente_leitura(path: Path) -> Iterator[InstantaneoArquivo]:
    """Confere hash e metadados antes/depois para detectar qualquer alteracao."""
    inicial = _instantaneo(path)
    try:
        yield inicial
    finally:
        final = _instantaneo(path)
        if final != inicial:
            raise ArquivoCorrompidoError(
                f"o arquivo foi alterado durante a leitura e nao pode ser aceito: {path}"
            )


def cabecalhos_originais(valores: Sequence[object]) -> list[str]:
    return ["" if value is None else str(value) for value in valores]


def cabecalhos_normalizados(valores: Sequence[object]) -> list[str]:
    return [normalize_text(value).comparavel if isinstance(value, str) else "" for value in valores]


def valor_obrigatorio(
    valores: Sequence[object],
    indices: dict[str, int],
    campo: str,
    *,
    arquivo: Path,
    aba: str,
    linha: int,
) -> object:
    value = valores[indices[campo]] if indices[campo] < len(valores) else None
    if value is None or (isinstance(value, str) and not value.strip()):
        raise LinhaIncompletaError(
            f"valor obrigatorio ausente para {campo}",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            coluna=get_column_letter(indices[campo] + 1),
            celula=f"{get_column_letter(indices[campo] + 1)}{linha}",
            cabecalho_esperado=campo,
            valor_recebido=value,
        )
    return value


def valor_opcional(valores: Sequence[object], indice: int | None) -> object | None:
    if indice is None or indice >= len(valores):
        return None
    value = valores[indice]
    return None if isinstance(value, str) and not value.strip() else value


def converter_moeda(
    value: object,
    *,
    arquivo: Path,
    aba: str,
    linha: int,
    coluna: int,
    cabecalho: str,
) -> Decimal:
    try:
        return parse_money(value)
    except MonetarioError as exc:
        letter = get_column_letter(coluna + 1)
        raise ValorMonetarioInvalidoError(
            f"valor monetario invalido para {cabecalho}",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            coluna=letter,
            celula=f"{letter}{linha}",
            cabecalho_esperado=cabecalho,
            valor_recebido=value,
        ) from exc


def converter_data(
    value: object,
    *,
    arquivo: Path,
    aba: str,
    linha: int,
    coluna: int,
    cabecalho: str,
) -> date:
    try:
        if isinstance(value, str) and " " in value.strip():
            value = value.strip().split(maxsplit=1)[0]
        return parse_date(value)
    except DataError as exc:
        letter = get_column_letter(coluna + 1)
        raise DataLeituraInvalidaError(
            f"data invalida para {cabecalho}",
            arquivo=arquivo,
            aba=aba,
            linha=linha,
            coluna=letter,
            celula=f"{letter}{linha}",
            cabecalho_esperado=cabecalho,
            valor_recebido=value,
        ) from exc


def converter_hora(value: object | None) -> time | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, str):
        for pattern in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(value.strip(), pattern).time()
            except ValueError:
                continue
    raise DataError(f"hora invalida: {value!r}")


def data_e_hora_venda(value: object, hora_explicita: object | None) -> tuple[date, time | None]:
    """Extrai hora de uma data combinada quando a coluna de hora nao existir."""
    if isinstance(value, datetime):
        return value.date(), converter_hora(hora_explicita) or value.time().replace(microsecond=0)
    if isinstance(value, str) and " " in value.strip():
        date_text, time_text = value.strip().split(maxsplit=1)
        return parse_date(date_text), converter_hora(hora_explicita) or converter_hora(time_text)
    return parse_date(value), converter_hora(hora_explicita)


def modalidade_por_texto(value: object) -> Modalidade | None:
    normalized = normalize_text(str(value)).comparavel
    if "pix" in normalized:
        return Modalidade.PIX
    if "debito" in normalized:
        return Modalidade.DEBITO
    if "credito" in normalized:
        return Modalidade.CREDITO
    return None


def linha_total_ou_vazia(valores: Sequence[object]) -> bool:
    non_empty = [value for value in valores if value is not None and str(value).strip()]
    if not non_empty:
        return True
    return any(normalize_text(str(value)).comparavel.startswith("total") for value in non_empty)


def dados_originais(
    *, headers: Sequence[str], values: Sequence[object], row_number: int
) -> tuple[dict[str, object], dict[str, str]]:
    raw_values: dict[str, object] = {}
    cell_map: dict[str, str] = {}
    for index, value in enumerate(values):
        header = (
            headers[index] if index < len(headers) and headers[index] else f"coluna_{index + 1}"
        )
        key = f"{header} ({get_column_letter(index + 1)})"
        raw_values[key] = value
        cell_map[key] = f"{get_column_letter(index + 1)}{row_number}"
    return {"valores": raw_values, "linha_original": row_number}, cell_map


def _instantaneo(path: Path) -> InstantaneoArquivo:
    try:
        stat = path.stat()
        return InstantaneoArquivo(
            tamanho_bytes=stat.st_size,
            modificado_em_ns=stat.st_mtime_ns,
            hash_sha256=sha256_file(path),
        )
    except OSError as exc:
        raise ArquivoCorrompidoError(
            f"nao foi possivel obter metadados do arquivo: {path}"
        ) from exc
