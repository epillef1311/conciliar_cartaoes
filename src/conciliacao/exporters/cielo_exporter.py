"""Exportador Excel do relatorio Cielo no padrao funcional LEO."""

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from conciliacao.domain.models import TransacaoOperadora
from conciliacao.matching.models import ResultadoConciliacaoOperadora
from conciliacao.processors.cielo_processor import CieloRelatorioProcessado
from conciliacao.utils.currency import quantize_money
from conciliacao.validators.common_validator import valor_original

CIELO_HEADERS = [
    "Data de pagamento",
    "Data da venda",
    "Hora da venda",
    "Tipo de lançamento",
    "Forma de pagamento",
    "Bandeira",
    "Valor bruto",
    "Taxa/tarifa",
    "Valor líquido",
    "Porcentagem",
    "Soma brutosoma",
    "soma liquido",
]

MATCHING_SHEET_NAME = "Conciliação Velo"
MATCHING_HEADERS = [
    "Origem",
    "Linha Cielo",
    "Data da venda",
    "Bandeira",
    "Valor Cielo",
    "Valor Velo",
    "Diferença",
    "Status",
    "Nível de confiança",
]

MONEY_FORMAT = r'\R\$\ * #,##0.00;[Red]\-\R\$\ * #,##0.00'
PERCENT_FORMAT = "0.00%"
DATE_FORMAT = "dd/mm/yyyy"
TIME_FORMAT = "hh:mm"


@dataclass(frozen=True, slots=True)
class CieloExportResult:
    caminho_saida: Path
    total_row: int
    ultima_linha_transacao: int
    subtotal_rows: tuple[int, ...]
    validacao_saida_ok: bool


class CieloExporter:
    def exportar(
        self,
        relatorio: CieloRelatorioProcessado,
        *,
        diretorio_saida: str | Path,
        sobrescrever: bool = False,
        matching: ResultadoConciliacaoOperadora | None = None,
    ) -> CieloExportResult:
        output_dir = Path(diretorio_saida)
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = _resolver_saida(
            output_dir
            / f"CIELO_CONCILIACAO_{relatorio.data_inicio.isoformat()}_A_"
            f"{relatorio.data_fim.isoformat()}.xlsx",
            sobrescrever=sobrescrever,
        )

        workbook = Workbook()
        worksheet = cast(Worksheet, workbook.active)
        if worksheet is None:
            raise ValueError("nao foi possivel criar a aba Planilha1")
        worksheet.title = "Planilha1"
        _escrever_planilha(worksheet, relatorio)
        if matching is not None:
            _escrever_comparacao_velo(workbook.create_sheet(MATCHING_SHEET_NAME), matching)
        workbook.save(output_path)
        workbook.close()

        result = CieloExportResult(
            caminho_saida=output_path,
            total_row=3 + len(relatorio.transacoes),
            ultima_linha_transacao=2 + len(relatorio.transacoes),
            subtotal_rows=tuple(
                3 + bloco.fim_indice for bloco in relatorio.blocos if bloco.tipo == "CARTAO"
            ),
            validacao_saida_ok=False,
        )
        _validar_exportacao(output_path, relatorio, result, matching)
        return CieloExportResult(
            caminho_saida=result.caminho_saida,
            total_row=result.total_row,
            ultima_linha_transacao=result.ultima_linha_transacao,
            subtotal_rows=result.subtotal_rows,
            validacao_saida_ok=True,
        )


def _escrever_planilha(worksheet: Worksheet, relatorio: CieloRelatorioProcessado) -> None:
    worksheet["A1"] = relatorio.titulo
    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 30
    for cell in worksheet[1]:
        cell.fill = _yellow_fill()
    for column, header in enumerate(CIELO_HEADERS, start=1):
        cell = worksheet.cell(2, column, header)
        cell.fill = _yellow_fill()
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()

    for index, transacao in enumerate(relatorio.transacoes, start=3):
        _escrever_transacao(worksheet, index, transacao)
    _escrever_subtotais(worksheet, relatorio)
    _escrever_total_geral(worksheet, relatorio)
    _aplicar_estilo(worksheet, relatorio, 3 + len(relatorio.transacoes))


def _escrever_comparacao_velo(
    worksheet: Worksheet, matching: ResultadoConciliacaoOperadora
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:I1")
    worksheet["A1"] = "CONCILIAÇÃO CIELO × VELO"
    worksheet["A1"].font = Font(bold=True, size=14, color="FFFFFFFF")
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 26

    worksheet.merge_cells("A2:I2")
    worksheet["A2"] = "Comparação gerada pelo workflow de conciliação Velo."
    worksheet["A2"].font = Font(italic=True, color="FF44546A")

    resumo_headers = [
        "Transações Cielo",
        "Registros Velo",
        "Conciliadas",
        "Somente Cielo",
        "Somente Velo",
        "Ambiguidades",
    ]
    resumo_values = [
        matching.resumo.transacoes_operadora,
        matching.resumo.registros_sistema,
        matching.resumo.conciliados,
        matching.resumo.nao_encontrados_no_sistema,
        matching.resumo.nao_encontrados_na_operadora,
        matching.resumo.ambiguidades,
    ]
    for summary_column, summary_header in enumerate(resumo_headers, start=1):
        cell = worksheet.cell(4, summary_column, summary_header)
        cell.fill = PatternFill(fill_type="solid", fgColor="FFD9EAF7")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _thin_border()
    for summary_column, summary_value in enumerate(resumo_values, start=1):
        cell = worksheet.cell(5, summary_column, summary_value)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _thin_border()

    for column, header in enumerate(MATCHING_HEADERS, start=1):
        cell = worksheet.cell(7, column, header)
        cell.fill = PatternFill(fill_type="solid", fgColor="FF5B9BD5")
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()

    for row, item in enumerate(matching.individuais, start=8):
        transacao = item.transacao_operadora
        registro = item.registro_sistema
        bandeira: str | int | None
        if transacao is None and registro is None:
            raise ValueError("resultado de matching sem transacao Cielo ou registro Velo")
        if transacao is not None:
            origem = "Cielo"
            linha_cielo: int | None = transacao.linha_original
            data = transacao.data_venda
            bandeira = transacao.bandeira
        else:
            assert registro is not None
            origem = "Somente Velo"
            linha_cielo = None
            data = registro.data_cadastro
            bandeira = registro.tipo_cartao
        values: list[Any] = [
            origem,
            linha_cielo,
            data,
            bandeira,
            item.valor_operadora,
            item.valor_sistema,
            item.diferenca,
            item.status.value.replace("_", " "),
            item.nivel_confianca.value.replace("_", " "),
        ]
        for value_column, value in enumerate(values, start=1):
            cell = worksheet.cell(row, value_column, value)
            cell.border = _thin_border()
            cell.alignment = Alignment(vertical="center")
        for money_column in (5, 6, 7):
            worksheet.cell(row, money_column).number_format = MONEY_FORMAT
        worksheet.cell(row, 3).number_format = DATE_FORMAT
        if item.status.value == "CONCILIADO":
            worksheet.cell(row, 8).fill = PatternFill(fill_type="solid", fgColor="FFE2F0D9")
        else:
            worksheet.cell(row, 8).fill = PatternFill(fill_type="solid", fgColor="FFFCE4D6")

    widths = {"A": 16, "B": 13, "C": 15, "D": 16, "E": 15, "F": 15, "G": 15, "H": 34, "I": 26}
    for column_letter, width in widths.items():
        worksheet.column_dimensions[column_letter].width = width
    worksheet.freeze_panes = "A8"


def _escrever_transacao(worksheet: Worksheet, row: int, transacao: TransacaoOperadora) -> None:
    values = [
        transacao.data_recebimento,
        transacao.data_venda,
        transacao.hora_venda,
        valor_original(transacao, "Tipo de lançamento"),
        valor_original(transacao, "Forma de pagamento"),
        transacao.bandeira,
        transacao.valor_bruto,
        transacao.taxa_original,
        transacao.valor_liquido,
        None if transacao.valor_bruto == Decimal("0.00") else f"=1-(I{row}/G{row})",
        None,
        None,
    ]
    for column, value in enumerate(values, start=1):
        worksheet.cell(row, column, value)


def _escrever_subtotais(worksheet: Worksheet, relatorio: CieloRelatorioProcessado) -> None:
    for bloco in relatorio.blocos:
        if bloco.tipo == "PIX":
            continue
        start_row = 3 + bloco.inicio_indice
        end_row = 3 + bloco.fim_indice
        worksheet.cell(end_row, 11, f"=SUM(G{start_row}:G{end_row})")
        worksheet.cell(end_row, 12, f"=SUM(I{start_row}:I{end_row})")


def _escrever_total_geral(worksheet: Worksheet, relatorio: CieloRelatorioProcessado) -> None:
    total_row = 3 + len(relatorio.transacoes)
    last_transaction = total_row - 1
    worksheet.cell(total_row, 7, f"=SUM(G3:G{last_transaction})")
    worksheet.cell(total_row, 8, f"=SUM(H3:H{last_transaction})")
    worksheet.cell(total_row, 9, f"=SUM(I3:I{last_transaction})")
    if relatorio.resumo.total_bruto != Decimal("0.00"):
        worksheet.cell(total_row, 10, f"=1-(I{total_row}/G{total_row})")


def _aplicar_estilo(
    worksheet: Worksheet, relatorio: CieloRelatorioProcessado, total_row: int
) -> None:
    worksheet["A1"].font = Font(bold=True, size=12)
    worksheet["A1"].fill = _yellow_fill()
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

    widths = {
        "A": 24,
        "B": 16,
        "C": 14,
        "D": 24,
        "E": 22,
        "F": 15,
        "G": 14,
        "H": 13,
        "I": 15,
        "J": 13,
        "K": 15,
        "L": 14,
    }
    for column_letter, width in widths.items():
        worksheet.column_dimensions[column_letter].width = width

    for cells in worksheet.iter_rows(min_row=3, max_row=total_row, min_col=1, max_col=12):
        for cell in cells:
            cell.border = _thin_border()
            cell.alignment = Alignment(vertical="center")
    for row_number in range(3, total_row + 1):
        worksheet.cell(row_number, 1).number_format = DATE_FORMAT
        worksheet.cell(row_number, 2).number_format = DATE_FORMAT
        worksheet.cell(row_number, 3).number_format = TIME_FORMAT
        for numeric_column in (7, 8, 9, 11, 12):
            worksheet.cell(row_number, numeric_column).number_format = MONEY_FORMAT
        worksheet.cell(row_number, 10).number_format = PERCENT_FORMAT
    for total_column in (7, 8, 9, 10, 11, 12):
        worksheet.cell(total_row, total_column).font = Font(bold=True)
        worksheet.cell(total_row, total_column).border = _total_border()
    _aplicar_cores_blocos(worksheet, relatorio)


def _aplicar_cores_blocos(worksheet: Worksheet, relatorio: CieloRelatorioProcessado) -> None:
    """Replica o padrao LEO: cartao por dia/modalidade e Pix em amarelo-claro."""
    dias_credito = {
        dia: indice
        for indice, dia in enumerate(
            sorted(
                {
                    relatorio.transacoes[bloco.inicio_indice].data_venda
                    for bloco in relatorio.blocos
                    if bloco.tipo == "CARTAO"
                    and relatorio.transacoes[bloco.inicio_indice].modalidade.name == "CREDITO"
                }
            )
        )
    }
    for bloco in relatorio.blocos:
        transacao = relatorio.transacoes[bloco.inicio_indice]
        color = _cor_bloco_cielo(transacao, dias_credito)
        fill = PatternFill(fill_type="solid", fgColor=f"FF{color}")
        for row in range(3 + bloco.inicio_indice, 4 + bloco.fim_indice):
            for column in range(1, 13):
                worksheet.cell(row, column).fill = fill


def _cor_bloco_cielo(transacao: TransacaoOperadora, dias_credito: dict[Any, int]) -> str:
    if transacao.modalidade.name == "PIX":
        return "FFFF99"
    if transacao.modalidade.name == "DEBITO":
        return "33CCCC"
    return "C0C0C0" if dias_credito[transacao.data_venda] % 2 == 0 else "FFFFFF"


def _validar_exportacao(
    path: Path,
    relatorio: CieloRelatorioProcessado,
    result: CieloExportResult,
    matching: ResultadoConciliacaoOperadora | None,
) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        expected_sheets = ["Planilha1"] if matching is None else ["Planilha1", MATCHING_SHEET_NAME]
        if workbook.sheetnames != expected_sheets:
            raise ValueError("abas do arquivo Cielo gerado nao coincidem com o layout esperado")
        worksheet = workbook["Planilha1"]
        headers = [worksheet.cell(2, column).value for column in range(1, 13)]
        if headers != CIELO_HEADERS:
            raise ValueError("cabecalhos Cielo gerados nao coincidem com o layout final")
        if worksheet.max_column != 12:
            raise ValueError("arquivo Cielo gerado deve conter exatamente 12 colunas")
        if result.ultima_linha_transacao != 2 + len(relatorio.transacoes):
            raise ValueError("quantidade de transacoes exportadas invalida")
        _validar_formulas(worksheet, relatorio, result)
        _validar_valores(worksheet, relatorio, result)
        if matching is not None:
            _validar_comparacao_velo(workbook[MATCHING_SHEET_NAME], matching)
        for sheet in workbook.worksheets:
            _validar_erros_formula(sheet)
    finally:
        workbook.close()


def _validar_formulas(
    worksheet: Worksheet, relatorio: CieloRelatorioProcessado, result: CieloExportResult
) -> None:
    for row, transacao in enumerate(relatorio.transacoes, start=3):
        expected = None if transacao.valor_bruto == Decimal("0.00") else f"=1-(I{row}/G{row})"
        if worksheet.cell(row, 10).value != expected:
            raise ValueError(f"formula de percentual invalida na linha {row}")
    for bloco in relatorio.blocos:
        if bloco.tipo == "PIX":
            continue
        start = 3 + bloco.inicio_indice
        end = 3 + bloco.fim_indice
        if worksheet.cell(end, 11).value != f"=SUM(G{start}:G{end})":
            raise ValueError(f"subtotal bruto invalido na linha {end}")
        if worksheet.cell(end, 12).value != f"=SUM(I{start}:I{end})":
            raise ValueError(f"subtotal liquido invalido na linha {end}")
    total = result.total_row
    last = result.ultima_linha_transacao
    expected_total = {
        (total, 7): f"=SUM(G3:G{last})",
        (total, 8): f"=SUM(H3:H{last})",
        (total, 9): f"=SUM(I3:I{last})",
        (total, 10): f"=1-(I{total}/G{total})",
    }
    for (row, col), formula in expected_total.items():
        if worksheet.cell(row, col).value != formula:
            raise ValueError(f"formula total invalida em {get_column_letter(col)}{row}")
    _validar_cores_blocos(worksheet, relatorio)


def _validar_cores_blocos(worksheet: Worksheet, relatorio: CieloRelatorioProcessado) -> None:
    dias_credito = {
        dia: indice
        for indice, dia in enumerate(
            sorted(
                {
                    relatorio.transacoes[bloco.inicio_indice].data_venda
                    for bloco in relatorio.blocos
                    if bloco.tipo == "CARTAO"
                    and relatorio.transacoes[bloco.inicio_indice].modalidade.name == "CREDITO"
                }
            )
        )
    }
    for bloco in relatorio.blocos:
        transacao = relatorio.transacoes[bloco.inicio_indice]
        expected = _cor_bloco_cielo(transacao, dias_credito)
        first_row = 3 + bloco.inicio_indice
        last_row = 3 + bloco.fim_indice
        for row in (first_row, last_row):
            actual = worksheet.cell(row, 1).fill.fgColor.rgb
            if actual != f"FF{expected}":
                raise ValueError(f"cor do bloco Cielo invalida na linha {row}")


def _validar_valores(
    worksheet: Worksheet, relatorio: CieloRelatorioProcessado, result: CieloExportResult
) -> None:
    bruto = Decimal("0.00")
    taxa = Decimal("0.00")
    liquido = Decimal("0.00")
    for row in range(3, result.ultima_linha_transacao + 1):
        bruto += Decimal(str(worksheet.cell(row, 7).value))
        taxa += Decimal(str(worksheet.cell(row, 8).value))
        liquido += Decimal(str(worksheet.cell(row, 9).value))
    if quantize_money(bruto) != relatorio.resumo.total_bruto:
        raise ValueError("total bruto exportado diverge do processamento")
    if quantize_money(taxa) != relatorio.resumo.total_taxa:
        raise ValueError("total taxa exportado diverge do processamento")
    if quantize_money(liquido) != relatorio.resumo.total_liquido:
        raise ValueError("total liquido exportado diverge do processamento")


def _validar_comparacao_velo(
    worksheet: Worksheet, matching: ResultadoConciliacaoOperadora
) -> None:
    headers = [worksheet.cell(7, column).value for column in range(1, 10)]
    if headers != MATCHING_HEADERS:
        raise ValueError("cabecalhos da comparacao Velo nao coincidem com o layout esperado")
    expected_summary = [
        matching.resumo.transacoes_operadora,
        matching.resumo.registros_sistema,
        matching.resumo.conciliados,
        matching.resumo.nao_encontrados_no_sistema,
        matching.resumo.nao_encontrados_na_operadora,
        matching.resumo.ambiguidades,
    ]
    summary = [worksheet.cell(5, column).value for column in range(1, 7)]
    if summary != expected_summary:
        raise ValueError("resumo da comparacao Velo diverge do matching")
    if worksheet.max_row != 7 + len(matching.individuais):
        raise ValueError("quantidade de linhas da comparacao Velo invalida")


def _validar_erros_formula(worksheet: Worksheet) -> None:
    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
    for row in worksheet.iter_rows():
        for cell in row:
            value = str(cell.value)
            if any(token in value for token in error_tokens):
                raise ValueError(f"formula com erro estrutural em {cell.coordinate}")


def _resolver_saida(path: Path, *, sobrescrever: bool) -> Path:
    if sobrescrever or not path.exists():
        return path
    suffix = datetime.now().strftime("%Y%m%d_%H%M%S")
    return path.with_name(f"{path.stem}_{suffix}{path.suffix}")


def _yellow_fill() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FFFFC000")


def _thin_border() -> Border:
    side = Side(style="thin", color="FF808080")
    return Border(left=side, right=side, top=side, bottom=side)


def _total_border() -> Border:
    side = Side(style="thin", color="FF000000")
    return Border(top=Side(style="medium", color="FF000000"), bottom=side)
