"""Leitura e geração atômica da planilha auxiliar de recebimentos QuickPay."""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from conciliacao.domain.enums import Modalidade
from conciliacao.utils.currency import parse_money
from conciliacao.utils.text import normalize_text

RECEBIMENTOS_BANCARIOS_SHEET = "Recebimentos Bancários"
RECEBIMENTOS_BANCARIOS_HEADERS = (
    "Data de recebimento",
    "Bandeira",
    "Modalidade",
    "Valor recebido no banco",
)
MONEY_FORMAT = r"\R\$\ * #,##0.00;[Red]\-\R\$\ * #,##0.00"


class RecebimentosBancariosError(ValueError):
    """Entrada bancária agregada inválida."""


@dataclass(frozen=True, slots=True)
class RecebimentoBancarioQuickPay:
    data_recebimento: date
    bandeira: str
    modalidade: Modalidade
    valor_recebido: Decimal
    linha_origem: int | None = None

    @property
    def chave(self) -> tuple[date, str, Modalidade]:
        return (
            self.data_recebimento,
            normalize_text(self.bandeira).comparavel,
            self.modalidade,
        )


def parse_recebimento_chat(raw: str) -> RecebimentoBancarioQuickPay:
    """Converte uma linha `data|bandeira|modalidade|valor` informada no chat."""
    parts = [item.strip() for item in raw.split("|")]
    if len(parts) != 4 or any(not item for item in parts):
        raise RecebimentosBancariosError("recebimento deve usar: data|bandeira|modalidade|valor")
    try:
        data_recebimento = date.fromisoformat(parts[0])
    except ValueError as exc:
        raise RecebimentosBancariosError(f"data de recebimento inválida: {parts[0]!r}") from exc
    modalidade = _modalidade(parts[2])
    try:
        valor = parse_money(parts[3])
    except ValueError as exc:
        raise RecebimentosBancariosError(f"valor bancário inválido: {parts[3]!r}") from exc
    if valor < Decimal("0.00"):
        raise RecebimentosBancariosError("valor recebido no banco não pode ser negativo")
    return RecebimentoBancarioQuickPay(data_recebimento, parts[1], modalidade, valor)


def criar_planilha_recebimentos_bancarios(
    recebimentos: list[RecebimentoBancarioQuickPay],
    *,
    diretorio_saida: Path,
    nome_arquivo: str = "RECEBIMENTOS_BANCARIOS_QUICKPAY.xlsx",
) -> Path:
    _validar_unicidade(recebimentos)
    diretorio_saida.mkdir(parents=True, exist_ok=True)
    destino = diretorio_saida / nome_arquivo
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = RECEBIMENTOS_BANCARIOS_SHEET
    sheet.append(list(RECEBIMENTOS_BANCARIOS_HEADERS))
    for recebimento in recebimentos:
        sheet.append(
            [
                recebimento.data_recebimento,
                recebimento.bandeira,
                recebimento.modalidade.value.title(),
                recebimento.valor_recebido,
            ]
        )
    for cell in sheet[1]:
        cell.fill = PatternFill(fill_type="solid", fgColor="FF5B9BD5")
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 1).number_format = "dd/mm/yyyy"
        sheet.cell(row, 4).number_format = MONEY_FORMAT
    for column, width in {"A": 20, "B": 18, "C": 16, "D": 26}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    with NamedTemporaryFile(dir=diretorio_saida, suffix=".xlsx", delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        workbook.save(temp_path)
        temp_path.replace(destino)
    finally:
        workbook.close()
        if temp_path.exists():
            temp_path.unlink()
    return destino


def ler_recebimentos_bancarios(path: Path) -> tuple[RecebimentoBancarioQuickPay, ...]:
    if not path.is_file():
        raise RecebimentosBancariosError(f"planilha bancária QuickPay não encontrada: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if RECEBIMENTOS_BANCARIOS_SHEET not in workbook.sheetnames:
            raise RecebimentosBancariosError(
                f"aba obrigatória ausente: {RECEBIMENTOS_BANCARIOS_SHEET}"
            )
        sheet = workbook[RECEBIMENTOS_BANCARIOS_SHEET]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        normalized = [normalize_text(str(value or "")).comparavel for value in headers]
        expected = [normalize_text(value).comparavel for value in RECEBIMENTOS_BANCARIOS_HEADERS]
        if normalized[: len(expected)] != expected:
            raise RecebimentosBancariosError("cabeçalhos da planilha bancária QuickPay inválidos")
        recebimentos = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = row[:4]
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            if any(value is None or str(value).strip() == "" for value in values):
                raise RecebimentosBancariosError(
                    "linha incompleta em "
                    f"{RECEBIMENTOS_BANCARIOS_SHEET}!A{row_number}:D{row_number}"
                )
            recebimentos.append(_from_cells(values, row_number))
    finally:
        workbook.close()
    if not recebimentos:
        raise RecebimentosBancariosError("planilha bancária QuickPay sem recebimentos")
    _validar_unicidade(recebimentos)
    return tuple(recebimentos)


def _from_cells(values: tuple[object, ...], row_number: int) -> RecebimentoBancarioQuickPay:
    raw_date, raw_brand, raw_modality, raw_value = values
    if isinstance(raw_date, datetime):
        data_recebimento = raw_date.date()
    elif isinstance(raw_date, date):
        data_recebimento = raw_date
    elif isinstance(raw_date, str):
        try:
            data_recebimento = date.fromisoformat(raw_date.strip())
        except ValueError as exc:
            raise RecebimentosBancariosError(
                f"data inválida em A{row_number}: {raw_date!r}"
            ) from exc
    else:
        raise RecebimentosBancariosError(f"data inválida em A{row_number}")
    modalidade = _modalidade(str(raw_modality))
    try:
        valor = parse_money(raw_value)
    except ValueError as exc:
        raise RecebimentosBancariosError(f"valor inválido em D{row_number}") from exc
    if valor < Decimal("0.00"):
        raise RecebimentosBancariosError(f"valor negativo em D{row_number}")
    return RecebimentoBancarioQuickPay(
        data_recebimento, str(raw_brand), modalidade, valor, row_number
    )


def _modalidade(value: str) -> Modalidade:
    comparable = normalize_text(value).comparavel
    if "credito" in comparable:
        return Modalidade.CREDITO
    if "debito" in comparable or "dedito" in comparable:
        return Modalidade.DEBITO
    if "pix" in comparable:
        return Modalidade.PIX
    raise RecebimentosBancariosError(f"modalidade inválida: {value!r}")


def _validar_unicidade(
    recebimentos: list[RecebimentoBancarioQuickPay] | tuple[RecebimentoBancarioQuickPay, ...],
) -> None:
    keys = [item.chave for item in recebimentos]
    if len(keys) != len(set(keys)):
        raise RecebimentosBancariosError(
            "há mais de um recebimento para a mesma data, bandeira e modalidade"
        )
